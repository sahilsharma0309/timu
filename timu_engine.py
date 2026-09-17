"""
Timu Engine — query-driven multi-site data extraction.
=======================================================

Core library behind the Timu agent. No web framework needed: you can use it
straight from Python or from the CLI. `timu_app.py` wraps this in an animated
browser UI.

    from timu_engine import TimuScraper, TimuConfig

    timu = TimuScraper(TimuConfig(max_pages_per_site=3))
    report = timu.run(["example.com", "https://other.org"], "find all emails and prices")
    print(report["common"])          # data shared across the sites
    print(report["sites"][0]["data"])  # per-site extraction

Design notes
------------
* Up to 10 targets per run (hard cap, matches the UI).
* A bare name like "wikipedia" is resolved by probing common TLDs.
* Query -> intents -> extractors. Intent keywords understand English and
  Hinglish ("keemat", "sampark", "tasveer", "sara data", ...).
* `css:` / `selector:` / `xpath-ish` hints in the query switch on a custom
  selector extractor.
* Optional LLM pass (`llm_fn`) turns the page text into query-shaped rows when
  you have an API key or an in-process model handle.
* robots.txt is honoured by default and politeness delay is per-host.
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
import urllib.parse as up
import urllib.robotparser as robotparser
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone
from typing import Any, Callable, Iterable

from bs4 import BeautifulSoup

try:                                     # only the legacy sequential fetcher needs requests
    import requests
    HAVE_REQUESTS = True
except Exception:                        # pragma: no cover - hardened hosts may block it
    requests = None                      # type: ignore[assignment]
    HAVE_REQUESTS = False

__all__ = ["TimuConfig", "TimuScraper", "PageResult", "export_json", "export_csv",
           "export_xlsx", "export_markdown", "MAX_TARGETS"]

MAX_TARGETS = 10
TIMU_VERSION = "2.0"

# ---------------------------------------------------------------- config ----


@dataclass
class TimuConfig:
    timeout: float = 20.0
    max_pages_per_site: int = 1          # 1 = only the given page; >1 crawls same-host links
    max_targets: int = MAX_TARGETS
    delay: float = 0.6                   # politeness delay between requests to one host
    respect_robots: bool = True
    concurrency: int = 5
    max_bytes: int = 4_000_000
    # --- v2 transport (timu_fetch) ---------------------------------------
    transport: str = "auto"              # auto = async httpx when available, else legacy requests
    render: str = "auto"                 # auto | http | browser  (browser needs Playwright)
    wait_for: str | None = None          # CSS selector the browser waits for before reading
    discover: bool = True                # sitemap.xml / feeds / rel=next for extra pages
    use_cache: bool = True               # on-disk cache + conditional GETs
    cache_ttl_s: int = 3600
    cache_path: str = ".timu_cache"
    total_concurrency: int = 16          # global in-flight request cap
    per_host_concurrency: int = 2        # in-flight cap for one host (politeness)
    user_agent: str = (                  # honest, contactable UA. Do not spoof a browser.
        f"TimuBot/{TIMU_VERSION} (+data-extraction agent; respects robots.txt)"
    )
    max_items: int = 400                 # cap per extracted field
    text_chars: int = 20000              # cap on stored page text
    llm_max_chars: int = 14000


# ------------------------------------------------------------- fetch layer ----

if HAVE_REQUESTS:
    _RequestsSSLError = requests.exceptions.SSLError
    _RequestsConnError = requests.exceptions.ConnectionError
    _RequestsTimeout = requests.exceptions.Timeout
else:                                    # placeholders so `except` clauses stay valid
    class _RequestsSSLError(Exception): ...
    class _RequestsConnError(Exception): ...
    class _RequestsTimeout(Exception): ...


@dataclass
class PageResult:
    target: str
    url: str
    final_url: str = ""
    status: int | None = None
    ok: bool = False
    html: str = ""
    text: str = ""
    elapsed_ms: int = 0
    error: str = ""
    content_type: str = ""
    blocked_by_robots: bool = False

    def brief(self) -> dict:
        d = asdict(self)
        d.pop("html", None)
        d["text"] = self.text[:400]
        return d


class Fetcher:
    """Session-pooled HTTP GET with robots.txt cache and per-host throttling."""

    def __init__(self, cfg: TimuConfig):
        if not HAVE_REQUESTS:
            raise RuntimeError("legacy fetcher needs `requests` (pip install requests); "
                               "the async transport in timu_fetch.py only needs httpx")
        self.cfg = cfg
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": cfg.user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en,hi;q=0.8",
        })
        self._robots: dict[str, robotparser.RobotFileParser | None] = {}
        self._last_hit: dict[str, float] = {}

    # -- robots ------------------------------------------------------------
    def _robots_for(self, url: str):
        host = up.urlsplit(url)._replace(path="", query="", fragment="").geturl()
        if host in self._robots:
            return self._robots[host]
        rp = robotparser.RobotFileParser()
        try:
            r = self.session.get(up.urljoin(host, "/robots.txt"), timeout=min(8, self.cfg.timeout))
            if r.status_code < 400 and r.text:
                rp.parse(r.text.splitlines())
            else:
                rp = None
        except Exception:
            rp = None
        self._robots[host] = rp
        return rp

    def allowed(self, url: str) -> bool:
        if not self.cfg.respect_robots:
            return True
        rp = self._robots_for(url)
        if rp is None:
            return True                     # no robots.txt served == no restriction
        try:
            return rp.can_fetch(self.cfg.user_agent, url)
        except Exception:
            return True

    # -- throttle ----------------------------------------------------------
    def _throttle(self, url: str):
        host = up.urlsplit(url).netloc
        last = self._last_hit.get(host)
        if last is not None:
            wait = self.cfg.delay - (time.time() - last)
            if wait > 0:
                time.sleep(wait)
        self._last_hit[host] = time.time()

    # -- get ---------------------------------------------------------------
    def get(self, url: str, target: str | None = None) -> PageResult:
        res = PageResult(target=target or url, url=url)
        if not self.allowed(url):
            res.error = "disallowed by robots.txt"
            res.blocked_by_robots = True
            return res
        self._throttle(url)
        t0 = time.time()
        try:
            r = self.session.get(url, timeout=self.cfg.timeout, stream=True,
                                 allow_redirects=True)
            res.status = r.status_code
            res.final_url = r.url
            res.content_type = (r.headers.get("Content-Type") or "").split(";")[0].strip()
            raw = r.raw.read(self.cfg.max_bytes, decode_content=True) or b""
            r.close()
            enc = r.encoding or "utf-8"
            body = raw.decode(enc, errors="replace")
            res.html = body
            res.ok = r.status_code < 400 and bool(body)
            if not res.ok and not res.error:
                res.error = f"HTTP {r.status_code}"
        except _RequestsSSLError as e:
            res.error = f"ssl error: {e.__class__.__name__}"
        except _RequestsConnError as e:
            msg = str(e)
            if "403" in msg and "CONNECT" in msg:
                res.error = ("blocked by network policy (sandbox allowlist) — "
                             "run Timu on your own machine or allow this domain")
            else:
                res.error = f"connection failed: {e.__class__.__name__}"
        except _RequestsTimeout:
            res.error = f"timeout after {self.cfg.timeout:g}s"
        except Exception as e:
            res.error = f"{e.__class__.__name__}: {e}"
        res.elapsed_ms = int((time.time() - t0) * 1000)
        return res


# ------------------------------------------------------- target resolution ----

_TLDS = (".com", ".org", ".net", ".io", ".in", ".co", ".ai", ".dev", ".gov", ".edu")


def normalise_target(raw: str) -> list[str]:
    """Turn user input into an ordered list of candidate URLs to try."""
    s = (raw or "").strip().strip('"\'' ).rstrip("/")
    if not s:
        return []
    if s.startswith(("http://", "https://")):
        return [s]
    if re.match(r"^[\w.-]+\.[a-z]{2,}(?:[/?#].*)?$", s, re.I):       # looks like a domain
        return [f"https://{s}", f"http://{s}"]
    slug = re.sub(r"[^a-z0-9-]+", "", s.lower().replace(" ", ""))     # bare name -> probe TLDs
    if not slug:
        return []
    cands = [f"https://{slug}{t}" for t in _TLDS]
    return cands + [f"https://www.{slug}.com"]


# --------------------------------------------------------------- parsing ----

_SOUP_PARSERS = ("lxml", "html.parser")


def make_soup(html: str) -> BeautifulSoup:
    for p in _SOUP_PARSERS:
        try:
            return BeautifulSoup(html, p)
        except Exception:
            continue
    return BeautifulSoup(html, "html.parser")


_SKIP_TEXT_IN = {"script", "style", "noscript", "template", "svg", "head", "meta", "link"}


def visible_text(soup: BeautifulSoup, limit: int = 20000) -> str:
    """Human-visible text. Non-destructive: the soup keeps its <script> blocks so
    JSON-LD / structured-data extraction still works after this call."""
    parts = []
    for node in soup.find_all(string=True):
        parent = getattr(node, "parent", None)
        if parent is not None and getattr(parent, "name", "") in _SKIP_TEXT_IN:
            continue
        t = str(node).strip()
        if t:
            parts.append(t)
    txt = re.sub(r"\n{2,}", "\n", "\n".join(parts))
    return txt[:limit]


# ------------------------------------------------------------ extractors ----

EMAIL_RE = re.compile(r"[\w.+-]+@(?:[\w-]+\.)+[A-Za-z]{2,24}")
PHONE_RE = re.compile(r"(?:(?:\+|00)\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?){2,4}\d{2,4}")
PRICE_RE = re.compile(
    r"(?:(?:Rs\.?|INR|₹|\$|USD|€|EUR|£|GBP|¥)\s?\d[\d,]*(?:\.\d{1,2})?)"
    r"|(?:\d[\d,]*(?:\.\d{1,2})?\s?(?:INR|USD|EUR|GBP|rupees|rs\.?))", re.I)
DATE_RE = re.compile(
    r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2}"
    r"|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4}"
    r"|\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})\b", re.I)
SOCIAL_HOSTS = ("facebook.com", "twitter.com", "x.com", "instagram.com", "linkedin.com",
                "youtube.com", "github.com", "t.me", "wa.me", "whatsapp.com", "pinterest.com",
                "tiktok.com", "reddit.com", "medium.com", "discord.gg", "threads.net")
STOPWORDS = set("""a an the of in on at for to from with and or is are was were be been being this
that these those all any its it's it as by about into over under more most find get list give show
me my mine we our us you your data info information page site website url nikalo nikal batao bata
kya konsa kaun sara saara sab poora pura do de dena chahiye karo kar please need want extract scrape
mujhe mera hai hain ka ki ke ko se mein me par aur ya bhi lao dhundo dhoondo""".split())


def _uniq(seq: Iterable[str], cap: int) -> list[str]:
    seen, out = set(), []
    for s in seq:
        s = (s or "").strip()
        if not s:
            continue
        k = s.lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(s)
        if len(out) >= cap:
            break
    return out


def _clean_phone(m: str) -> str | None:
    digits = re.sub(r"\D", "", m)
    if not 7 <= len(digits) <= 15:
        return None
    if re.fullmatch(r"(\d)\1+", digits):
        return None
    return m.strip()


def x_meta(soup, text, url, cfg) -> dict:
    g = {}
    if soup.title and soup.title.string:
        g["title"] = soup.title.string.strip()
    for name, key in (("description", "description"), ("keywords", "keywords"),
                      ("author", "author"), ("og:title", "og_title"),
                      ("og:description", "og_description"), ("og:site_name", "site_name"),
                      ("og:image", "og_image")):
        tag = soup.find("meta", attrs={"name": name}) or soup.find("meta", property=name)
        if tag and tag.get("content"):
            g[key] = tag["content"].strip()[:500]
    h1 = soup.find("h1")
    if h1:
        g["h1"] = h1.get_text(" ", strip=True)[:300]
    lang = soup.find("html")
    if lang and lang.get("lang"):
        g["lang"] = lang["lang"]
    return g


def _hrefs_with_scheme(soup, scheme: str) -> list[str]:
    """href values for one scheme, without needing soupsieve (CSS engine) installed."""
    return [a["href"].strip() for a in soup.find_all("a", href=True)
            if a["href"].strip().lower().startswith(scheme)]


def x_emails(soup, text, url, cfg) -> list[str]:
    found = set(EMAIL_RE.findall(text))
    for h in _hrefs_with_scheme(soup, "mailto:"):
        found.add(up.unquote(h[7:].split("?")[0]))
    bad = ("example.com", "domain.com", "email.com", "sentry.io", ".png", ".jpg", ".webp")
    return _uniq((e for e in found if not any(b in e.lower() for b in bad)), cfg.max_items)


def x_phones(soup, text, url, cfg) -> list[str]:
    found = []
    for h in _hrefs_with_scheme(soup, "tel:"):
        found.append(up.unquote(h[4:]))
    for m in PHONE_RE.findall(text):
        c = _clean_phone(m)
        if c:
            found.append(re.sub(r"\s{2,}", " ", c))
    return _uniq(found, cfg.max_items)


def x_links(soup, text, url, cfg) -> list[dict]:
    base = url
    out, seen = [], set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href.startswith(("javascript:", "#", "mailto:", "tel:")):
            continue
        absu = up.urljoin(base, href)
        if absu in seen:
            continue
        seen.add(absu)
        out.append({"text": a.get_text(" ", strip=True)[:160] or None, "url": absu,
                    "internal": up.urlsplit(absu).netloc == up.urlsplit(base).netloc})
        if len(out) >= cfg.max_items:
            break
    return out


def x_social(soup, text, url, cfg) -> list[str]:
    out = []
    for a in soup.find_all("a", href=True):
        u = up.urljoin(url, a["href"])
        host = up.urlsplit(u).netloc.lower().removeprefix("www.")
        if any(host == h or host.endswith("." + h) for h in SOCIAL_HOSTS):
            out.append(u.split("?")[0])
    return _uniq(out, cfg.max_items)


def x_images(soup, text, url, cfg) -> list[dict]:
    out = []
    for img in soup.find_all("img"):
        src = img.get("src") or img.get("data-src") or img.get("data-original")
        if not src or src.startswith("data:"):
            continue
        out.append({"url": up.urljoin(url, src.strip()),
                    "alt": (img.get("alt") or "").strip()[:200] or None})
        if len(out) >= cfg.max_items:
            break
    return out


def x_headings(soup, text, url, cfg) -> list[dict]:
    out = []
    for lvl in ("h1", "h2", "h3", "h4"):
        for h in soup.find_all(lvl):
            t = h.get_text(" ", strip=True)
            if t:
                out.append({"level": lvl, "text": t[:300]})
    return out[:cfg.max_items]


def x_tables(soup, text, url, cfg) -> list[dict]:
    tables = []
    for ti, tb in enumerate(soup.find_all("table")):
        rows = []
        for tr in tb.find_all("tr"):
            cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
            if any(cells):
                rows.append(cells)
        if len(rows) < 2:
            continue
        header, body = rows[0], rows[1:]
        if len(set(header)) < len(header) or not all(header):
            header = [f"col{i+1}" for i in range(max(len(r) for r in rows))]
            body = rows
        recs = [{header[i] if i < len(header) else f"col{i+1}": v for i, v in enumerate(r)}
                for r in body]
        cap = tb.find("caption")
        tables.append({"index": ti, "caption": cap.get_text(" ", strip=True) if cap else None,
                       "columns": header, "rows": recs[:cfg.max_items]})
        if len(tables) >= 25:
            break
    return tables


def x_prices(soup, text, url, cfg) -> list[str]:
    return _uniq((re.sub(r"\s+", " ", p).strip() for p in PRICE_RE.findall(text)), cfg.max_items)


def x_dates(soup, text, url, cfg) -> list[str]:
    return _uniq(DATE_RE.findall(text), cfg.max_items)


def x_structured(soup, text, url, cfg) -> list[Any]:
    out = []
    for s in soup.find_all("script", attrs={"type": re.compile("ld\\+json", re.I)}):
        raw = s.string or s.get_text() or ""
        try:
            out.append(json.loads(raw))
        except Exception:
            try:
                out.append(json.loads(re.sub(r",\s*([}\]])", r"\1", raw)))
            except Exception:
                continue
    return out[:20]


def x_paragraphs(soup, text, url, cfg) -> list[str]:
    ps = [p.get_text(" ", strip=True) for p in soup.find_all("p")]
    return [p for p in ps if len(p) > 40][:cfg.max_items]


def x_fulltext(soup, text, url, cfg) -> str:
    return text


def x_items(soup, text, url, cfg) -> list[dict]:
    """Generic repeated-card detector: listings, products, articles, search results."""
    best, best_score = [], 0
    for parent in soup.find_all(["ul", "ol", "div", "section", "main", "tbody"]):
        kids = [c for c in parent.find_all(recursive=False) if getattr(c, "name", None)]
        if len(kids) < 3:
            continue
        sig = {}
        for k in kids:
            key = (k.name, " ".join(sorted(k.get("class") or []))[:80])
            sig.setdefault(key, []).append(k)
        for key, group in sig.items():
            if len(group) < 3:
                continue
            texts = [g.get_text(" ", strip=True) for g in group]
            filled = [t for t in texts if len(t) > 15]
            score = len(filled) * min(sum(len(t) for t in filled) / max(len(filled), 1), 400)
            if score > best_score:
                best_score, best = score, group
    out = []
    for node in best[:cfg.max_items]:
        t = re.sub(r"\s+", " ", node.get_text(" ", strip=True))
        if len(t) < 10:
            continue
        a = node.find("a", href=True)
        img = node.find("img")
        price = PRICE_RE.search(t)
        head = node.find(["h1", "h2", "h3", "h4", "strong", "b"])
        out.append({
            "name": (head.get_text(" ", strip=True) if head else t)[:200],
            "text": t[:600],
            "link": up.urljoin(url, a["href"]) if a else None,
            "price": re.sub(r"\s+", " ", price.group(0)).strip() if price else None,
            "image": up.urljoin(url, img.get("src")) if img and img.get("src") else None,
        })
    return out


def x_addresses(soup, text, url, cfg) -> list[str]:
    hits = []
    for tag in soup.find_all(["address"]):
        hits.append(tag.get_text(" ", strip=True))
    pat = re.compile(r"[^\n]{0,80}(?:street|st\.|road|rd\.|avenue|ave\.|lane|marg|nagar|sector"
                     r"|block|floor|building|plot|pincode|pin code|zip)[^\n]{0,80}", re.I)
    hits += pat.findall(text)
    pin = re.compile(r"[^\n]{0,60}\b\d{6}\b[^\n]{0,20}")
    hits += pin.findall(text)[:20]
    return _uniq((re.sub(r"\s+", " ", h).strip() for h in hits), 60)


EXTRACTORS: dict[str, Callable] = {
    "meta": x_meta, "emails": x_emails, "phones": x_phones, "links": x_links,
    "social_links": x_social, "images": x_images, "headings": x_headings,
    "tables": x_tables, "prices": x_prices, "dates": x_dates,
    "structured_data": x_structured, "paragraphs": x_paragraphs,
    "items": x_items, "addresses": x_addresses, "full_text": x_fulltext,
}

# intent keyword map (English + Hinglish)
INTENT_WORDS: dict[str, tuple[str, ...]] = {
    "emails": ("email", "e-mail", "mail", "gmail", "contact", "sampark", "id"),
    "phones": ("phone", "mobile", "number", "tel", "telephone", "whatsapp", "contact",
               "sampark", "call", "numbers", "no."),
    "prices": ("price", "pricing", "cost", "rate", "keemat", "kimat", "daam", "dam",
               "amount", "fee", "fees", "salary", "budget", "mrp", "discount", "offer"),
    "tables": ("table", "tables", "tabular", "spreadsheet", "row", "rows", "column",
               "columns", "stats", "statistics", "data table", "chart"),
    "links": ("link", "links", "url", "urls", "href", "hyperlink", "sitemap", "pages"),
    "social_links": ("social", "facebook", "twitter", "instagram", "linkedin", "youtube",
                     "handle", "profile", "github"),
    "images": ("image", "images", "photo", "photos", "picture", "pictures", "tasveer",
               "thumbnail", "logo", "img"),
    "headings": ("heading", "headings", "title", "titles", "topic", "topics", "section",
                 "sections", "headline", "headlines", "outline", "menu"),
    "items": ("product", "products", "item", "items", "listing", "listings", "catalog",
              "catalogue", "saman", "card", "cards", "job", "jobs", "course", "courses",
              "post", "posts", "article", "articles", "news", "blog", "result", "results",
              "review", "reviews", "property", "flat", "car", "book", "books", "movie",
              "recipe", "event", "events", "offer", "deal", "deals", "vacancy"),
    "dates": ("date", "dates", "tarikh", "deadline", "published", "time", "schedule",
              "year", "when", "kab"),
    "addresses": ("address", "addresses", "location", "pata", "office", "branch", "pincode",
                  "city", "where", "kahan", "map"),
    "paragraphs": ("paragraph", "paragraphs", "description", "about", "summary", "content",
                   "body", "detail", "details", "vivaran", "jankari", "info"),
    "full_text": ("everything", "all data", "sara data", "saara data", "sab kuch", "full text",
                  "whole page", "complete", "raw text", "dump", "poora", "pura data"),
    "structured_data": ("schema", "json-ld", "jsonld", "structured", "metadata", "seo",
                        "rich snippet"),
}
ALWAYS_ON = ("meta", "structured_data")
DEFAULT_INTENTS = ("meta", "headings", "items", "links", "structured_data", "query_matches")


def detect_intents(query: str) -> list[str]:
    q = f" {(query or '').lower()} "
    hits = []
    for field_name, words in INTENT_WORDS.items():
        for w in words:
            if (" " in w and w in q) or re.search(rf"\b{re.escape(w)}\b", q):
                hits.append(field_name)
                break
    if not hits:
        hits = list(DEFAULT_INTENTS)
    for a in ALWAYS_ON:
        if a not in hits:
            hits.append(a)
    if "query_matches" not in hits:
        hits.append("query_matches")
    order = list(EXTRACTORS) + ["query_matches", "css_matches", "llm_rows"]
    return sorted(set(hits), key=lambda f: order.index(f) if f in order else 99)


def query_terms(query: str) -> list[str]:
    toks = re.findall(r"[a-zA-Z\u0900-\u097F][\w'-]{2,}", query or "")
    return [t for t in dict.fromkeys(t.lower() for t in toks) if t not in STOPWORDS][:12]


def x_query_matches(soup, text, url, cfg, terms: list[str]) -> list[dict]:
    if not terms:
        return []
    out = []
    for line in text.split("\n"):
        s = line.strip()
        if len(s) < 3 or len(s) > 600:
            continue
        low = s.lower()
        hit = [t for t in terms if t in low]
        if hit:
            out.append({"match": s, "terms": hit})
    out.sort(key=lambda d: -len(d["terms"]))
    return out[:cfg.max_items]


CSS_HINT = re.compile(r"(?:css|selector|select)\s*[:=]\s*(.+?)(?:$|\||;)", re.I)


def css_from_query(query: str) -> str | None:
    m = CSS_HINT.search(query or "")
    return m.group(1).strip().strip('"\'') if m else None


def x_css_matches(soup, sel: str, url: str, cfg: TimuConfig) -> list[dict]:
    try:
        nodes = soup.select(sel)
    except Exception as e:
        return [{"error": f"bad selector: {e}"}]
    out = []
    for n in nodes[:cfg.max_items]:
        out.append({"text": n.get_text(" ", strip=True)[:500],
                    "href": up.urljoin(url, n["href"]) if n.has_attr("href") else None,
                    "src": up.urljoin(url, n["src"]) if n.has_attr("src") else None})
    return out


# ------------------------------------------------------------- LLM shaping ----

LLM_SYSTEM = (
    "You are Timu's extraction head. You receive text scraped from ONE web page and a user "
    "query. Return ONLY a JSON object: {\"answer\": short string, \"rows\": [ {..}, .. ], "
    "\"fields\": [names]}. 'rows' must be flat objects with consistent keys holding the data "
    "the query asks for, taken verbatim from the page. If the page does not contain it, return "
    "empty rows and say so in 'answer'. No prose outside the JSON."
)


def llm_rows(llm_fn: Callable[[str], str], query: str, page_text: str, url: str,
             cfg: TimuConfig) -> dict:
    prompt = (f"USER QUERY:\n{query}\n\nPAGE URL: {url}\n\nPAGE TEXT (truncated):\n"
              f"{page_text[:cfg.llm_max_chars]}\n\nReturn the JSON object now.")
    try:
        raw = llm_fn(prompt)
    except Exception as e:
        return {"error": f"llm call failed: {e.__class__.__name__}: {e}"}
    if not isinstance(raw, str):
        raw = str(raw)
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        return {"error": "llm returned no JSON", "raw": raw[:500]}
    try:
        return json.loads(m.group(0))
    except Exception:
        return {"error": "llm JSON parse failed", "raw": raw[:500]}


def anthropic_llm(api_key: str, model: str = "claude-sonnet-4-5") -> Callable[[str], str]:
    """Build an llm_fn backed by the Anthropic Messages API (stdlib urllib, no extra deps)."""
    import urllib.request

    def _call(prompt: str) -> str:
        payload = json.dumps({"model": model, "max_tokens": 4096, "system": LLM_SYSTEM,
                              "messages": [{"role": "user", "content": prompt}]}).encode()
        req = urllib.request.Request(
            "https://api.anthropic.com/v1/messages", data=payload, method="POST",
            headers={"x-api-key": api_key, "anthropic-version": "2023-06-01",
                     "content-type": "application/json"})
        with urllib.request.urlopen(req, timeout=120) as r:
            blocks = json.loads(r.read().decode("utf-8")).get("content", [])
        return "".join(b.get("text", "") for b in blocks if b.get("type") == "text")
    return _call


# ------------------------------------------------------------- the scraper ----


@dataclass
class TimuScraper:
    cfg: TimuConfig = field(default_factory=TimuConfig)
    llm_fn: Callable[[str], str] | None = None

    # -- one site ----------------------------------------------------------
    def _resolve(self, target: str, fetcher: Fetcher) -> PageResult:
        cands = normalise_target(target)
        if not cands:
            return PageResult(target=target, url=target, error="empty or unusable target")
        last = None
        for c in cands:
            res = fetcher.get(c, target=target)
            if res.ok:
                return res
            last = res
            if res.blocked_by_robots or (res.status and res.status < 500):
                if len(cands) == 1:
                    return res
        return last or PageResult(target=target, url=cands[0], error="unreachable")

    def _crawl_extra(self, first: PageResult, fetcher: Fetcher) -> list[PageResult]:
        if self.cfg.max_pages_per_site <= 1 or not first.ok:
            return []
        soup = make_soup(first.html)
        base = first.final_url or first.url
        host = up.urlsplit(base).netloc
        seen = {base.rstrip("/")}
        queue = []
        for a in soup.find_all("a", href=True):
            u = up.urljoin(base, a["href"]).split("#")[0].rstrip("/")
            if up.urlsplit(u).netloc != host or u in seen:
                continue
            if re.search(r"\.(?:pdf|zip|jpg|jpeg|png|gif|svg|mp4|mp3|docx?|xlsx?)$", u, re.I):
                continue
            seen.add(u)
            queue.append(u)
        out = []
        for u in queue[: self.cfg.max_pages_per_site - 1]:
            r = fetcher.get(u, target=first.target)
            if r.ok:
                out.append(r)
        return out

    def _extract_page(self, page: PageResult, query: str, intents: list[str],
                      sel: str | None) -> dict:
        soup = make_soup(page.html)
        text = visible_text(soup, self.cfg.text_chars)
        page.text = text
        data: dict[str, Any] = {}
        terms = query_terms(query)
        for name in intents:
            try:
                if name == "query_matches":
                    v = x_query_matches(soup, text, page.final_url or page.url, self.cfg, terms)
                elif name in EXTRACTORS:
                    v = EXTRACTORS[name](soup, text, page.final_url or page.url, self.cfg)
                else:
                    continue
            except Exception as e:
                v = {"error": f"{e.__class__.__name__}: {e}"}
            if v:
                data[name] = v
        if sel:
            data["css_matches"] = x_css_matches(soup, sel, page.final_url or page.url, self.cfg)
        if self.llm_fn:
            data["llm_rows"] = llm_rows(self.llm_fn, query, text, page.final_url or page.url,
                                        self.cfg)
        return data

    # -- v2: one concurrent fetch pass for every target, then extract -------
    def _prefetch_all(self, targets: list[str],
                      progress: Callable[[dict], None] | None) -> dict | None:
        """Fetch every target (and its extra pages) through timu_fetch.

        Returns {target: {"pages": [PageResult], "meta": {...}}} or None when the
        async transport is unavailable, in which case the legacy path is used.
        """
        if self.cfg.transport == "legacy":
            return None
        try:
            import timu_fetch as tfetch
        except Exception:
            return None
        if not tfetch.HAVE_HTTPX:
            return None
        say = progress or (lambda _e: None)
        if self.cfg.render == "browser" and not tfetch.HAVE_PLAYWRIGHT:
            say({"event": "warn", "message": "Playwright not installed — falling back to HTTP "
                                             "render (pip install playwright && "
                                             "python -m playwright install chromium)"})
        fcfg = tfetch.FetchConfig(
            timeout=self.cfg.timeout, respect_robots=self.cfg.respect_robots,
            per_host_delay=self.cfg.delay, per_host_concurrency=self.cfg.per_host_concurrency,
            total_concurrency=self.cfg.total_concurrency, use_cache=self.cfg.use_cache,
            cache_ttl_s=self.cfg.cache_ttl_s, cache_path=self.cfg.cache_path,
            user_agent=self.cfg.user_agent, max_bytes=self.cfg.max_bytes)
        cand = {t: normalise_target(t) for t in targets}
        try:
            batches = tfetch.run_sync(tfetch.fetch_sites(
                targets, cfg=fcfg, pages=self.cfg.max_pages_per_site, render=self.cfg.render,
                discover=self.cfg.discover, wait_for=self.cfg.wait_for,
                candidates_map=cand, on_event=say))
        except Exception as e:                                        # noqa: BLE001
            say({"event": "warn", "message": f"async transport failed ({e.__class__.__name__}: "
                                             f"{e}); using legacy fetcher"})
            return None
        out: dict[str, dict] = {}
        for b in batches:
            t = b.get("target") or b.get("url")
            pages: list[PageResult] = []
            for r in b["all_results"]:
                pages.append(PageResult(
                    target=t, url=r.url, final_url=r.final_url or r.url, status=r.status,
                    ok=r.ok, html=r.body, elapsed_ms=r.elapsed_ms, error=r.error,
                    content_type=r.content_type, blocked_by_robots=r.blocked_by_robots))
            disc = b.get("discovery") or {}
            out[t] = {"pages": pages, "meta": {
                "renderer": next((r.renderer for r in b["all_results"] if r.ok), "none"),
                "from_cache": sum(1 for r in b["all_results"] if r.from_cache),
                "attempts": max([r.attempts for r in b["all_results"]] or [1]),
                "discovery": {k: v for k, v in disc.items() if k in ("counts", "sitemaps",
                                                                     "feeds", "pagination")},
                "api_endpoints": b.get("api_endpoints") or [],
                "notes": b.get("notes") or []}}
        return out

    def _run_site(self, target: str, query: str, intents: list[str], sel: str | None,
                  progress: Callable[[dict], None] | None,
                  pre: dict | None = None) -> dict:
        say = progress or (lambda _e: None)
        if pre is not None:
            pages = pre["pages"]
            first = next((p for p in pages if p.ok), pages[0] if pages else
                         PageResult(target=target, url=target, error="nothing fetched"))
            meta = pre.get("meta") or {}
        else:
            fetcher = Fetcher(self.cfg)
            say({"event": "site_start", "target": target})
            first = self._resolve(target, fetcher)
            pages = [first] + (self._crawl_extra(first, fetcher) if first.ok else [])
            meta = {"renderer": "http-legacy"}
        site = {"target": target, "url": first.final_url or first.url, "ok": first.ok,
                "status": first.status, "error": first.error, "elapsed_ms": first.elapsed_ms,
                "fetch": meta,
                "pages_fetched": sum(1 for p in pages if p.ok), "data": {}, "page_urls": []}
        if not first.ok:
            say({"event": "site_done", "target": target, "ok": False, "error": first.error})
            return site
        merged: dict[str, Any] = {}
        for p in [x for x in pages if x.ok]:
            site["page_urls"].append(p.final_url or p.url)
            d = self._extract_page(p, query, intents, sel)
            for k, v in d.items():
                if isinstance(v, list):
                    merged.setdefault(k, [])
                    if isinstance(merged[k], list):
                        merged[k].extend(v)
                elif isinstance(v, dict):
                    merged.setdefault(k, {})
                    if isinstance(merged[k], dict):
                        for kk, vv in v.items():
                            merged[k].setdefault(kk, vv)
                elif isinstance(v, str):
                    merged[k] = (merged.get(k, "") + ("\n\n" if merged.get(k) else "") + v)[
                        : self.cfg.text_chars * 2]
        for k, v in list(merged.items()):
            if isinstance(v, list) and v and isinstance(v[0], (str, int, float)):
                merged[k] = _uniq((str(x) for x in v), self.cfg.max_items)
            elif isinstance(v, list):
                out, seen = [], set()
                for item in v:
                    key = json.dumps(item, sort_keys=True, default=str)[:400]
                    if key in seen:
                        continue
                    seen.add(key)
                    out.append(item)
                merged[k] = out[: self.cfg.max_items]
        site["data"] = merged
        site["counts"] = {k: (len(v) if isinstance(v, (list, dict)) else len(str(v)))
                          for k, v in merged.items()}
        say({"event": "site_done", "target": target, "ok": True,
             "url": site["url"], "counts": site["counts"]})
        return site

    # -- many sites --------------------------------------------------------
    def run(self, targets: list[str], query: str = "", *,
            progress: Callable[[dict], None] | None = None) -> dict:
        targets = [t.strip() for t in targets if t and t.strip()][: self.cfg.max_targets]
        if not targets:
            raise ValueError("no targets given")
        intents = detect_intents(query)
        sel = css_from_query(query)
        t0 = time.time()
        say = progress or (lambda _e: None)
        say({"event": "run_start", "targets": targets, "intents": intents,
             "selector": sel, "llm": bool(self.llm_fn)})
        pre_map = self._prefetch_all(targets, progress)
        say({"event": "transport", "mode": "async-http2" if pre_map else "legacy-requests",
             "render": self.cfg.render, "cache": self.cfg.use_cache})
        sites: list[dict] = [None] * len(targets)                     # type: ignore
        workers = max(1, min(self.cfg.concurrency, len(targets)))
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futs = {pool.submit(self._run_site, t, query, intents, sel, progress,
                                (pre_map or {}).get(t)): i
                    for i, t in enumerate(targets)}
            for f in as_completed(futs):
                i = futs[f]
                try:
                    sites[i] = f.result()
                except Exception as e:
                    sites[i] = {"target": targets[i], "ok": False,
                                "error": f"{e.__class__.__name__}: {e}", "data": {}}
        report = {
            "timu_version": TIMU_VERSION,
            "query": query,
            "intents": intents,
            "selector": sel,
            "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "elapsed_s": round(time.time() - t0, 2),
            "targets": targets,
            "sites": sites,
            "summary": {
                "requested": len(targets),
                "succeeded": sum(1 for s in sites if s and s.get("ok")),
                "failed": [{"target": s["target"], "error": s.get("error")}
                           for s in sites if s and not s.get("ok")],
                "total_records": sum(sum(v for v in (s.get("counts") or {}).values())
                                     for s in sites if s),
                "pages_fetched": sum((s.get("pages_fetched") or 0) for s in sites if s),
                "transport": "async-http2" if pre_map else "legacy-requests",
                "renderers": sorted({(s.get("fetch") or {}).get("renderer", "none")
                                     for s in sites if s}),
                "cache_hits": sum((s.get("fetch") or {}).get("from_cache", 0)
                                  for s in sites if s),
                "api_endpoints_found": sum(len((s.get("fetch") or {}).get("api_endpoints") or [])
                                           for s in sites if s),
                "fetch_notes": [n for s in sites if s
                                for n in ((s.get("fetch") or {}).get("notes") or [])],
            },
            "common": find_common(sites) if len([s for s in sites if s and s.get("ok")]) > 1 else
                      {"note": "common-data analysis needs 2+ successful sites",
                       "shared_values": {}, "shared_fields": [], "overlap_matrix": {}},
        }
        say({"event": "run_done", "summary": report["summary"]})
        return report


# ------------------------------------------------------ common-data finder ----

def _flat_values(v: Any) -> list[str]:
    """Flatten any extracted structure into comparable strings (booleans are noise)."""
    out = []
    if isinstance(v, bool):
        return out
    if isinstance(v, str):
        out.append(v)
    elif isinstance(v, dict):
        for x in v.values():
            out += _flat_values(x)
    elif isinstance(v, list):
        for x in v:
            out += _flat_values(x)
    elif v is not None:
        out.append(str(v))
    return out


def _norm(s: str) -> str:
    s = re.sub(r"\s+", " ", (s or "")).strip().lower()
    s = re.sub(r"^https?://(www\.)?", "", s).rstrip("/")
    return s


def find_common(sites: list[dict]) -> dict:
    live = [s for s in sites if s and s.get("ok") and s.get("data")]
    if len(live) < 2:
        return {"shared_values": {}, "shared_fields": [], "overlap_matrix": {},
                "note": "need 2+ successful sites"}
    fields = [set(s["data"].keys()) for s in live]
    shared_fields = sorted(set.intersection(*fields)) if fields else []

    per_site_vals: list[dict[str, dict[str, str]]] = []
    for s in live:
        m: dict[str, dict[str, str]] = {}
        for f, v in s["data"].items():
            d = {}
            for raw in _flat_values(v):
                n = _norm(raw)
                if 2 < len(n) <= 300:
                    d.setdefault(n, raw if isinstance(raw, str) else str(raw))
            m[f] = d
        per_site_vals.append(m)

    shared_values: dict[str, list[dict]] = {}
    for f in shared_fields:
        keysets = [set(m.get(f, {})) for m in per_site_vals]
        inter = set.intersection(*keysets) if keysets else set()
        rows = []
        for k in sorted(inter):
            rows.append({"value": per_site_vals[0][f].get(k, k), "field": f,
                         "sites": len(live)})
        if rows:
            shared_values[f] = rows[:300]

    # value seen on >=2 sites, across any field (loose but useful)
    allvals: dict[str, set[int]] = {}
    display: dict[str, str] = {}
    for i, m in enumerate(per_site_vals):
        for f, d in m.items():
            for k, orig in d.items():
                if len(k) < 4:
                    continue
                allvals.setdefault(k, set()).add(i)
                display.setdefault(k, orig)
    multi = [{"value": display[k], "site_count": len(v),
              "sites": sorted(live[i]["target"] for i in v)}
             for k, v in allvals.items() if len(v) > 1]
    multi.sort(key=lambda r: (-r["site_count"], len(r["value"])))

    overlap = {}
    for i in range(len(live)):
        for j in range(i + 1, len(live)):
            a = {k for m in [per_site_vals[i]] for d in m.values() for k in d if len(k) >= 4}
            b = {k for m in [per_site_vals[j]] for d in m.values() for k in d if len(k) >= 4}
            if a or b:
                overlap[f"{live[i]['target']} ∩ {live[j]['target']}"] = {
                    "shared": len(a & b), "jaccard": round(len(a & b) / max(len(a | b), 1), 4)}
    return {
        "sites_compared": [s["target"] for s in live],
        "shared_fields": shared_fields,
        "shared_values": shared_values,
        "values_on_multiple_sites": multi[:400],
        "overlap_matrix": overlap,
    }


# ---------------------------------------------------------------- exports ----

def _rows_for_table(report: dict) -> list[dict]:
    rows = []
    for s in report.get("sites", []):
        if not s:
            continue
        if not s.get("ok"):
            rows.append({"site": s.get("target"), "url": s.get("url"), "field": "_error",
                         "key": "", "value": s.get("error")})
            continue
        for f, v in (s.get("data") or {}).items():
            if isinstance(v, dict):
                for k, vv in v.items():
                    rows.append({"site": s["target"], "url": s.get("url"), "field": f,
                                 "key": str(k), "value": _stringify(vv)})
            elif isinstance(v, list):
                for i, item in enumerate(v):
                    if isinstance(item, dict):
                        for k, vv in item.items():
                            rows.append({"site": s["target"], "url": s.get("url"),
                                         "field": f, "key": f"{i}.{k}",
                                         "value": _stringify(vv)})
                    else:
                        rows.append({"site": s["target"], "url": s.get("url"), "field": f,
                                     "key": str(i), "value": _stringify(item)})
            else:
                rows.append({"site": s["target"], "url": s.get("url"), "field": f,
                             "key": "", "value": _stringify(v)})
    for f, vals in (report.get("common", {}).get("shared_values") or {}).items():
        for r in vals:
            rows.append({"site": "_COMMON_", "url": "", "field": f, "key": "shared",
                         "value": _stringify(r.get("value"))})
    return rows


def _stringify(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return str(v)
    return json.dumps(v, ensure_ascii=False, default=str)[:2000]


def export_json(report: dict) -> str:
    return json.dumps(report, ensure_ascii=False, indent=2, default=str)


def export_csv(report: dict) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=["site", "url", "field", "key", "value"],
                       extrasaction="ignore", lineterminator="\n")
    w.writeheader()
    for r in _rows_for_table(report):
        w.writerow(r)
    return buf.getvalue()


def export_markdown(report: dict) -> str:
    L = [f"# Timu extraction report", "",
         f"- **Query:** {report.get('query') or '(none)'}",
         f"- **Generated:** {report.get('generated_at')}  ",
         f"- **Fields targeted:** {', '.join(report.get('intents', []))}",
         f"- **Sites:** {report['summary']['succeeded']}/{report['summary']['requested']} ok",
         f"- **Records:** {report['summary']['total_records']}", ""]
    for s in report.get("sites", []):
        if not s:
            continue
        L.append(f"## {s.get('target')}")
        if not s.get("ok"):
            L += [f"FAILED — {s.get('error')}", ""]
            continue
        L.append(f"`{s.get('url')}` · pages: {s.get('pages_fetched')}")
        for f, v in (s.get("data") or {}).items():
            n = len(v) if isinstance(v, (list, dict)) else 1
            L.append(f"\n### {f} ({n})")
            if isinstance(v, dict):
                L += [f"- **{k}**: {_stringify(vv)[:300]}" for k, vv in list(v.items())[:40]]
            elif isinstance(v, list):
                L += [f"- {_stringify(x)[:300]}" for x in v[:60]]
                if len(v) > 60:
                    L.append(f"- … {len(v)-60} more")
            else:
                L.append(_stringify(v)[:4000])
        L.append("")
    c = report.get("common") or {}
    L += ["## Common across sites", ""]
    if c.get("shared_values"):
        for f, rows in c["shared_values"].items():
            L.append(f"### {f} — {len(rows)} shared")
            L += [f"- {_stringify(r['value'])[:250]}" for r in rows[:40]]
    else:
        L.append(c.get("note", "No exact values shared by every site."))
    if c.get("values_on_multiple_sites"):
        L += ["", "### Values seen on 2+ sites", ""]
        L += [f"- {_stringify(r['value'])[:200]}  _({r['site_count']} sites)_"
              for r in c["values_on_multiple_sites"][:60]]
    return "\n".join(L)


def export_xlsx(report: dict, path: str) -> str:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "extraction"
    ws.append(["site", "url", "field", "key", "value"])
    for r in _rows_for_table(report):
        ws.append([r["site"], r["url"], r["field"], r["key"], str(r["value"])[:32000]])
    ws2 = wb.create_sheet("common")
    ws2.append(["field", "value", "sites"])
    for f, rows in (report.get("common", {}).get("shared_values") or {}).items():
        for r in rows:
            ws2.append([f, _stringify(r["value"])[:32000], r.get("sites")])
    for r in (report.get("common", {}).get("values_on_multiple_sites") or [])[:500]:
        ws2.append(["_any_", _stringify(r["value"])[:32000], ", ".join(r["sites"])])
    ws3 = wb.create_sheet("summary")
    ws3.append(["key", "value"])
    for k in ("query", "generated_at", "elapsed_s", "timu_version"):
        ws3.append([k, str(report.get(k))])
    for k, v in report["summary"].items():
        ws3.append([k, _stringify(v)])
    for sheet in (ws, ws2, ws3):
        for col, width in zip("ABCDE", (28, 46, 18, 22, 90)):
            sheet.column_dimensions[col].width = width
    wb.save(path)
    return path


# -------------------------------------------------------------------- CLI ----

def _cli(argv: list[str] | None = None) -> int:
    import argparse
    p = argparse.ArgumentParser(prog="timu", description="Timu — query-driven web data agent")
    p.add_argument("targets", nargs="+", help="URLs or site names (max 10)")
    p.add_argument("-q", "--query", default="", help="what to extract")
    p.add_argument("-o", "--out", default="timu_output", help="output basename")
    p.add_argument("--pages", type=int, default=1, help="pages per site (crawl depth)")
    p.add_argument("--no-robots", action="store_true", help="ignore robots.txt (your call)")
    p.add_argument("--llm-key", default=None, help="Anthropic API key for LLM shaping")
    p.add_argument("--formats", default="json,csv,md", help="json,csv,md,xlsx")
    p.add_argument("--render", default="auto", choices=["auto", "http", "browser"],
                   help="auto escalates to Playwright only when the HTML looks JS-built")
    p.add_argument("--wait-for", default=None, help="CSS selector the browser waits for")
    p.add_argument("--transport", default="auto", choices=["auto", "legacy"],
                   help="auto = async HTTP/2 via timu_fetch; legacy = sequential requests")
    p.add_argument("--no-cache", action="store_true", help="skip the on-disk page cache")
    p.add_argument("--no-discover", action="store_true",
                   help="do not use sitemap.xml / feeds / rel=next to find extra pages")
    a = p.parse_args(argv)
    cfg = TimuConfig(max_pages_per_site=max(1, a.pages), respect_robots=not a.no_robots,
                     render=a.render, wait_for=a.wait_for, transport=a.transport,
                     use_cache=not a.no_cache, discover=not a.no_discover)
    llm = anthropic_llm(a.llm_key) if a.llm_key else None
    timu = TimuScraper(cfg=cfg, llm_fn=llm)
    rep = timu.run(a.targets, a.query,
                   progress=lambda e: print(f"[timu] {e.get('event')}: "
                                            f"{e.get('target') or e.get('summary') or ''}"))
    fmts = {f.strip().lower() for f in a.formats.split(",")}
    if "json" in fmts:
        open(f"{a.out}.json", "w", encoding="utf-8").write(export_json(rep))
    if "csv" in fmts:
        open(f"{a.out}.csv", "w", encoding="utf-8", newline="").write(export_csv(rep))
    if "md" in fmts:
        open(f"{a.out}.md", "w", encoding="utf-8").write(export_markdown(rep))
    if "xlsx" in fmts:
        export_xlsx(rep, f"{a.out}.xlsx")
    s = rep["summary"]
    print(f"[timu] done — {s['succeeded']}/{s['requested']} sites, "
          f"{s['total_records']} records -> {a.out}.*")
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())
